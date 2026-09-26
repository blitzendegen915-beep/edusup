"""管理職向け会計報告書のテスト。python3 tests/test_admin_report.py で実行する。

壊した設定は一時ディレクトリのコピーで試し、実データは書き換えない。
"""
import datetime
import re
import shutil
import sys
import tempfile
from pathlib import Path

import yaml

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from scripts.kaikei import admin_report, model  # noqa: E402

CAMP = "2026-summer"
passed = failed = 0


def check(name, cond):
    global passed, failed
    passed, failed = (passed + 1, failed) if cond else (passed, failed + 1)
    print(("OK   " if cond else "FAIL ") + name)


def with_config(mutate):
    """report-admin の設定だけを書き換えた一時データで build_figures を実行する。"""
    tmp = Path(tempfile.mkdtemp())
    orig = model.DATA_DIR
    try:
        shutil.copytree(orig, tmp / "rugby", ignore=shutil.ignore_patterns("receipts"))
        p = tmp / "rugby" / f"report-admin-{CAMP}.yml"
        cfg = yaml.safe_load(p.read_text())
        mutate(cfg)
        p.write_text(yaml.safe_dump(cfg, allow_unicode=True))
        model.DATA_DIR = tmp / "rugby"
        try:
            admin_report.build_figures(CAMP)
            return None
        except SystemExit as e:
            return str(e)
    finally:
        model.DATA_DIR = orig
        shutil.rmtree(tmp)


# ── 実データ ───────────────────────────────────
f = admin_report.build_figures(CAMP)
camp = model.load(CAMP)
check("収入合計は会計システムの収入と一致", f["income_total"] == model.balance(camp)["income_total"])
check("支出合計は支出の各行の合計と一致", f["expense_total"] == sum(r[1] for r in f["expense_rows"]))
check("差引は 収入−支出", f["diff"] == f["income_total"] - f["expense_total"])
reported = sum(r[1] for r in f["expense_rows"]
               if not r[0].startswith(("宿泊・食事", "バス代", "予備費")))
check("台帳の支出は、報告書の費目と合宿会計外に過不足なく振り分けられている",
      reported + f["outside_total"] == sum(e.amount for e in camp.expenses))
check("ホテルは請求書額で計上される",
      any(r[1] == camp.hotel_invoice_total for r in f["expense_rows"]))
check("見積・精算中の行は『確定』扱いにしない",
      all(not r[2].startswith("確定") for r in f["expense_rows"] if r[0].startswith(("バス代", "雑費"))))

# ── 設定の不整合で止まるか ───────────────────────
def drop_r004(cfg):
    for ids in cfg["expense_lines"].values():
        if "r004" in ids:
            ids.remove("r004")
msg = with_config(drop_r004)
check("台帳の行を設定から外すとエラーで止まる（報告書から黙って抜けない）",
      msg is not None and "r004" in msg and "抜け落ちる" in msg)

def dup_r004(cfg):
    cfg["outside_camp"]["対戦校への手土産"].append("r004")
msg = with_config(dup_r004)
check("同じ領収書を2つの費目に入れるとエラーで止まる", msg is not None and "複数" in msg)

def ghost(cfg):
    next(iter(cfg["expense_lines"].values())).append("r999")
msg = with_config(ghost)
check("台帳に無い番号を書くとエラーで止まる", msg is not None and "r999" in msg)

# ── 出力文書 ───────────────────────────────────
out = Path(tempfile.mkdtemp()) / "t.docx"
admin_report.write(CAMP, out)
from docx import Document  # noqa: E402
d = Document(out)
text = "\n".join([p.text for p in d.paragraphs]
                 + [c.text for t in d.tables for r in t.rows for c in r.cells])
W = "月火水木金土日"
bad = [(m, dd, w) for m, dd, w in re.findall(r"(\d+)月(\d+)日（(.)）", text)
       if W[datetime.date(camp.start.year, int(m), int(dd)).weekday()] != w]
check("曜日の表記がすべて正しい", not bad)
flat = text.replace(" ", "").replace("　", "")
leaked = [p.name for p in camp.roster if p.name.replace(" ", "").replace("　", "") in flat]
check("生徒の氏名が報告書に載っていない", not leaked)
check("口座番号が報告書に載っていない", not re.search(r"普通\s*\d{7}", text))
check("領収書の無い支出は『領収書なし』と明記される", "領収書なし" in text)
shutil.rmtree(out.parent)

print(f"\n{passed} passed, {failed} failed")
sys.exit(1 if failed else 0)
