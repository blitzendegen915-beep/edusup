"""
注文集計。

データの流れ:
  order-<id>-responses.csv  フォーム回答の原文（手を加えない）
  order-<id>-lines.csv      備考欄まで読んで解釈した注文明細（人が判断して書く）
  order-<id>.yml            品名・サイズの並び・未回答を洗う対象学年
  roster-<年度>.csv         部員名簿（氏名・学年・組の正）

備考欄の自由記述（「追加でLを1枚」等）を機械に解釈させると誤読するので、
解釈は人が lines.csv に書き、ここでは原文との突き合わせと集計だけを行う。
"""
import argparse
import csv
import sys
import unicodedata
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

import yaml

ROOT = Path(__file__).resolve().parents[2]
DATA_DIR = ROOT / "data" / "rugby"
OUTPUT_DIR = ROOT / "output"

YELLOW = "FFFF00"


def norm_name(s: str) -> str:
    """全角/半角スペースや全角英数の揺れを吸収した突き合わせ用の氏名。"""
    s = unicodedata.normalize("NFKC", s or "")
    return "".join(s.split())


@dataclass
class Response:
    ts: datetime
    answer: str
    note: str
    klass: str
    name: str


@dataclass
class Line:
    name: str
    size: str
    qty: int
    status: str
    basis: str


def load(order_id: str, roster_year: str = "2026"):
    cfg = yaml.safe_load((DATA_DIR / f"order-{order_id}.yml").read_text())
    with open(DATA_DIR / f"order-{order_id}-responses.csv", newline="") as f:
        rows = list(csv.reader(f))[1:]
    responses = [
        Response(datetime.strptime(r[0], "%Y/%m/%d %H:%M:%S"), r[1].strip(),
                 r[2].strip(), r[3].strip(), r[4].strip())
        for r in rows if any(c.strip() for c in r)
    ]
    with open(DATA_DIR / f"order-{order_id}-lines.csv", newline="") as f:
        lines = [Line(r["name"].strip(), r["size"].strip(), int(r["qty"]),
                      r["status"].strip(), r["basis"].strip())
                 for r in csv.DictReader(f)]
    with open(DATA_DIR / f"roster-{roster_year}.csv", newline="") as f:
        roster = list(csv.DictReader(f))
    return cfg, responses, lines, roster


def latest_by_person(responses):
    """同一人物の複数回答は最新のタイムスタンプを採用する（rugby-order の規約）。"""
    by = defaultdict(list)
    for r in responses:
        by[norm_name(r.name)].append(r)
    latest = {k: max(v, key=lambda r: r.ts) for k, v in by.items()}
    dupes = {k: sorted(v, key=lambda r: r.ts) for k, v in by.items() if len(v) > 1}
    return latest, dupes


def analyze(order_id: str):
    cfg, responses, lines, roster = load(order_id)
    roster_by = {norm_name(r["name"]): r for r in roster}
    latest, dupes = latest_by_person(responses)
    declines = set(cfg.get("decline_values", []))
    sizes = cfg["sizes"]

    errors, warns, infos = [], [], []

    # 1. 明細の氏名が名簿にあるか
    for ln in lines:
        if norm_name(ln.name) not in roster_by:
            errors.append(f"明細の氏名「{ln.name}」が名簿にありません（表記揺れか誤記）")
    # 2. 回答者が名簿にあるか
    for k, r in latest.items():
        if k not in roster_by:
            errors.append(f"回答者「{r.name}」が名簿にありません（表記揺れか誤記）")
    # 3. サイズが定義にあるか
    for ln in lines:
        if ln.size not in sizes:
            errors.append(f"{ln.name}: サイズ「{ln.size}」が定義 {sizes} にありません")

    lines_by = defaultdict(list)
    for ln in lines:
        lines_by[norm_name(ln.name)].append(ln)

    # 4. 購入する回答者には明細があり、購入しない回答者には明細がないか
    for k, r in latest.items():
        has = k in lines_by
        if r.answer in declines and has:
            errors.append(f"{r.name}: 最新回答は「{r.answer}」なのに明細があります")
        if r.answer not in declines and not has:
            errors.append(f"{r.name}: 最新回答は「{r.answer}」なのに明細がありません（転記漏れ）")
    # 5. 回答していない人の明細が紛れていないか
    for k in lines_by:
        if k not in latest:
            errors.append(f"{lines_by[k][0].name}: 回答が無いのに明細があります")
    # 6. 最新回答の回答欄のサイズが明細に含まれているか（転記ミスの検出）
    for k, r in latest.items():
        if r.answer in declines or k not in lines_by:
            continue
        got = {ln.size for ln in lines_by[k]}
        if r.answer not in got:
            warns.append(f"{r.name}: 最新回答の回答欄は「{r.answer}」だが明細のサイズは {sorted(got)}。解釈を再確認")

    # 7. 複数回答
    for k, rs in dupes.items():
        hist = " → ".join(f"{r.ts:%m/%d %H:%M} {r.answer}" for r in rs)
        infos.append(f"{rs[-1].name}: 回答{len(rs)}回（{hist}）。最新を採用")

    # 8. 未回答者（対象学年）
    target = {str(g) for g in cfg.get("target_grades", [])}
    non_resp = [r for r in roster
                if r["grade"] in target and r.get("attends", "yes") != "__never__"
                and norm_name(r["name"]) not in latest]

    # 9. 対象学年外からの回答
    outside = [r for k, r in latest.items()
               if k in roster_by and roster_by[k]["grade"] not in target]

    pending = [ln for ln in lines if ln.status != "確定"]

    def tally(ls):
        c = Counter()
        for ln in ls:
            c[ln.size] += ln.qty
        return {s: c.get(s, 0) for s in sizes}

    return dict(
        cfg=cfg, responses=responses, lines=lines, roster_by=roster_by,
        latest=latest, dupes=dupes, errors=errors, warns=warns, infos=infos,
        non_resp=non_resp, outside=outside, pending=pending,
        total_all=tally(lines),
        total_confirmed=tally([ln for ln in lines if ln.status == "確定"]),
    )


def fmt_tally(t):
    parts = [f"{s}:{n}" for s, n in t.items() if n]
    return "  ".join(parts) + f"  ／ 計 {sum(t.values())}枚"


def cmd_check(order_id: str) -> int:
    a = analyze(order_id)
    cfg = a["cfg"]
    buyers = {norm_name(ln.name) for ln in a["lines"]}
    print(f"=== {cfg['item']} 注文チェック ===\n")
    print(f"回答 {len(a['responses'])}件 ／ 回答者 {len(a['latest'])}名 ／ 購入 {len(buyers)}名\n")

    print("--- サイズ別集計 ---")
    print(f"  全明細        : {fmt_tally(a['total_all'])}")
    print(f"  確定分のみ    : {fmt_tally(a['total_confirmed'])}\n")

    if a["errors"]:
        print(f"--- ❌ エラー（{len(a['errors'])}件）---")
        for m in a["errors"]:
            print(f"❌ {m}")
        print()
    if a["pending"]:
        print(f"--- ⚠️ 発注前に確認が要るもの（{len(a['pending'])}件）---")
        for ln in a["pending"]:
            print(f"⚠️ {ln.name}  {ln.size}×{ln.qty}\n     {ln.basis}")
        print()
    if a["warns"]:
        print(f"--- ⚠️ 要確認（{len(a['warns'])}件）---")
        for m in a["warns"]:
            print(f"⚠️ {m}")
        print()
    grades = "・".join(f"{g}年" for g in cfg.get("target_grades", []))
    print(f"--- 未回答（対象: {grades}）{len(a['non_resp'])}名 ---")
    for r in a["non_resp"]:
        print(f"  {r['grade']}年{r['class']}組 {r['name']}")
    if a["outside"]:
        print(f"\n--- 対象学年外からの回答 {len(a['outside'])}名 ---")
        for r in a["outside"]:
            g = a["roster_by"][norm_name(r.name)]
            print(f"  {g['grade']}年{g['class']}組 {r.name}（回答: {r.answer}）")
    if a["infos"]:
        print("\n--- 参考 ---")
        for m in a["infos"]:
            print(f"  {m}")
    print()
    if a["errors"]:
        print("❌ エラーがあります。lines.csv を直してから注文票を作ってください。")
        return 1
    print("✅ 回答原文と注文明細の突き合わせに矛盾はありません。")
    return 0


def _write_sheets(order_id: str, confirmed_only: bool):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    a = analyze(order_id)
    if a["errors"]:
        raise SystemExit("❌ check でエラーが出ています。先に lines.csv を直してください。")
    cfg = a["cfg"]
    rb = a["roster_by"]
    thin = Side(style="thin", color="999999")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_fill = PatternFill("solid", fgColor="D9E2F3")
    yellow = PatternFill("solid", fgColor=YELLOW)
    bold = Font(bold=True)

    def header(ws, cols, widths):
        for i, (c, w) in enumerate(zip(cols, widths), 1):
            cell = ws.cell(row=1, column=i, value=c)
            cell.font, cell.fill, cell.border = bold, head_fill, border
            ws.column_dimensions[cell.column_letter].width = w

    def sort_key(ln):
        g = rb[norm_name(ln.name)]
        return (int(g["grade"]), g["class"], norm_name(ln.name), cfg["sizes"].index(ln.size))

    send_lines = [ln for ln in a["lines"] if ln.status == "確定" or not confirmed_only]
    send_lines.sort(key=sort_key)
    OUTPUT_DIR.mkdir(exist_ok=True)

    # ── 業者送付用 ──────────────────────────────
    wb = Workbook()
    ws = wb.active
    ws.title = "サイズ別集計"
    ws["A1"], ws["A1"].font = cfg["item"], Font(bold=True, size=14)
    header_row = 3
    for i, c in enumerate(["サイズ", "数量"], 1):
        cell = ws.cell(row=header_row, column=i, value=c)
        cell.font, cell.fill, cell.border = bold, head_fill, border
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 12, 10
    r = header_row + 1
    first = r
    for s in cfg["sizes"]:
        ws.cell(row=r, column=1, value=s).border = border
        ws.cell(row=r, column=2, value=f'=SUMIF(注文一覧!C:C,"{s}",注文一覧!D:D)').border = border
        r += 1
    ws.cell(row=r, column=1, value="合計").font = bold
    tot = ws.cell(row=r, column=2, value=f"=SUM(B{first}:B{r-1})")
    tot.font, tot.border = bold, border

    ws2 = wb.create_sheet("注文一覧")
    header(ws2, ["学年・組", "生徒名", "サイズ", "数量", "品名"], [10, 16, 8, 8, 30])
    for i, ln in enumerate(send_lines, 2):
        g = rb[norm_name(ln.name)]
        for j, v in enumerate([f"{g['grade']}年{g['class']}組", g["name"], ln.size, ln.qty, cfg["item"]], 1):
            ws2.cell(row=i, column=j, value=v).border = border
    vendor_path = OUTPUT_DIR / f"order-{order_id}_業者送付用.xlsx"
    wb.save(vendor_path)

    # ── 確認用（部内向け。業者には送らない） ─────────
    wb = Workbook()
    ws = wb.active
    ws.title = "全明細と根拠"
    header(ws, ["学年・組", "生徒名", "サイズ", "数量", "状態", "解釈の根拠"], [10, 16, 8, 8, 8, 70])
    for i, ln in enumerate(sorted(a["lines"], key=sort_key), 2):
        g = rb[norm_name(ln.name)]
        vals = [f"{g['grade']}年{g['class']}組", g["name"], ln.size, ln.qty, ln.status, ln.basis]
        for j, v in enumerate(vals, 1):
            cell = ws.cell(row=i, column=j, value=v)
            cell.border = border
            cell.alignment = Alignment(wrap_text=(j == 6), vertical="top")
            if ln.status != "確定":
                cell.fill = yellow

    ws = wb.create_sheet("未回答")
    header(ws, ["学年・組", "生徒名", "状態"], [10, 16, 10])
    for i, rr in enumerate(a["non_resp"], 2):
        for j, v in enumerate([f"{rr['grade']}年{rr['class']}組", rr["name"], "未確認"], 1):
            cell = ws.cell(row=i, column=j, value=v)
            cell.border, cell.fill = border, yellow

    ws = wb.create_sheet("回答原文")
    header(ws, ["タイムスタンプ", "回答", "備考", "クラス（回答）", "氏名（回答）", "採用"], [18, 10, 50, 12, 14, 8])
    for i, rr in enumerate(sorted(a["responses"], key=lambda x: x.ts), 2):
        used = a["latest"][norm_name(rr.name)] is rr
        vals = [rr.ts.strftime("%Y/%m/%d %H:%M:%S"), rr.answer, rr.note, rr.klass, rr.name,
                "採用" if used else "破棄（古い回答）"]
        for j, v in enumerate(vals, 1):
            cell = ws.cell(row=i, column=j, value=v)
            cell.border = border
            cell.alignment = Alignment(wrap_text=(j == 3), vertical="top")
    review_path = OUTPUT_DIR / f"order-{order_id}_確認用.xlsx"
    wb.save(review_path)
    return vendor_path, review_path, send_lines, a


def cmd_sheet(order_id: str, confirmed_only: bool) -> int:
    vendor, review, send_lines, a = _write_sheets(order_id, confirmed_only)
    c = Counter()
    for ln in send_lines:
        c[ln.size] += ln.qty
    t = {s: c.get(s, 0) for s in a["cfg"]["sizes"]}
    print(f"✅ 業者送付用: {vendor}")
    print(f"   {fmt_tally(t)}（{'確定分のみ' if confirmed_only else '要確認分も含む'}）")
    print(f"✅ 部内確認用: {review}")
    if a["pending"] and not confirmed_only:
        print(f"\n⚠️ 要確認の {len(a['pending'])}件を解釈どおりに含めています。"
              "確認して変わったら lines.csv を直して再生成してください。")
    return 0


def main(argv=None):
    p = argparse.ArgumentParser(prog="python3 -m scripts.chumon",
                                description="ウェア・用具の注文集計")
    sub = p.add_subparsers(dest="cmd", required=True)
    c = sub.add_parser("check", help="回答原文と注文明細を突き合わせ、集計する")
    c.add_argument("order_id")
    s = sub.add_parser("sheet", help="業者送付用と部内確認用のExcelを作る")
    s.add_argument("order_id")
    s.add_argument("--confirmed-only", action="store_true",
                   help="状態が「確定」の明細だけで業者送付用を作る")
    args = p.parse_args(argv)
    if args.cmd == "check":
        sys.exit(cmd_check(args.order_id))
    sys.exit(cmd_sheet(args.order_id, args.confirmed_only))
